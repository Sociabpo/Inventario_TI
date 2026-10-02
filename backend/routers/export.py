from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse, Response
from pydantic import BaseModel
from typing import List, Optional, Any
from datetime import datetime
from html import escape
import io
import re

from routers.auth import get_current_user

router = APIRouter(prefix="/api/export", tags=["Exportación"])

# ── Topes de tamaño (exportación construye Excel/PDF → amplifica memoria) ──────
MAX_BODY_BYTES = 20 * 1024 * 1024   # 20 MB; rechazo temprano vía Content-Length
MAX_EXPORT_FILAS = 100_000          # muy por encima de cualquier export real


def _limitar_body(request: Request):
    """Rechazo barato y temprano por Content-Length (puede faltar/falsearse)."""
    cl = request.headers.get("content-length")
    if cl and cl.isdigit() and int(cl) > MAX_BODY_BYTES:
        raise HTTPException(status_code=413, detail="La solicitud es demasiado grande")


def _validar_filas(data: "ExportRequest"):
    """Garantía a prueba de spoofing: tope de filas ya parseadas."""
    if len(data.filas) > MAX_EXPORT_FILAS:
        raise HTTPException(
            status_code=413,
            detail=f"El export supera el máximo de {MAX_EXPORT_FILAS} filas",
        )


# ── Schemas ──────────────────────────────────────────────
class ColumnaExport(BaseModel):
    key: str
    label: str

class ExportRequest(BaseModel):
    titulo: str = "Reporte"
    subtitulo: Optional[str] = None
    columnas: List[ColumnaExport]
    filas: List[dict]


# ── Helpers ──────────────────────────────────────────────
def _fmt(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Sí" if v else "No"
    return str(v)

def _slug(texto: str) -> str:
    s = re.sub(r"[^\w\-]+", "_", texto.strip(), flags=re.UNICODE)
    return s.strip("_") or "reporte"

def _nombre_archivo(titulo: str, ext: str) -> str:
    return f"{_slug(titulo)}_{datetime.now().strftime('%Y-%m-%d')}.{ext}"


# ── Excel ────────────────────────────────────────────────
@router.post("/excel")
def export_excel(data: ExportRequest, current_user = Depends(get_current_user), _=Depends(_limitar_body)):
    _validar_filas(data)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte"

    ncols = max(1, len(data.columnas))
    thin = Side(style="thin", color="D9E1EC")
    borde = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Título
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=data.titulo)
    c.font = Font(bold=True, size=14, color="0E1C34")
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 22

    # Subtítulo / metadatos
    sub = data.subtitulo or f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')} — {len(data.filas)} registro(s)"
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    cs = ws.cell(row=2, column=1, value=sub)
    cs.font = Font(size=9, color="6A7E96", italic=True)

    # Encabezados (fila 4)
    header_row = 4
    header_fill = PatternFill("solid", fgColor="0E1C34")
    header_font = Font(bold=True, color="FFFFFF", size=10)
    for j, col in enumerate(data.columnas, start=1):
        cell = ws.cell(row=header_row, column=j, value=col.label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = borde

    # Datos
    anchos = [len(col.label) for col in data.columnas]
    for i, fila in enumerate(data.filas, start=header_row + 1):
        for j, col in enumerate(data.columnas, start=1):
            txt = _fmt(fila.get(col.key))
            cell = ws.cell(row=i, column=j, value=txt)
            cell.font = Font(size=10, color="1F2D3D")
            cell.border = borde
            cell.alignment = Alignment(vertical="center")
            if len(txt) > anchos[j - 1]:
                anchos[j - 1] = len(txt)

    # Ancho de columnas (con tope razonable)
    for j, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(j)].width = min(max(ancho + 2, 10), 50)

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{_nombre_archivo(data.titulo, "xlsx")}"'},
    )


# ── PDF ──────────────────────────────────────────────────
def _build_html(data: ExportRequest) -> str:
    sub = data.subtitulo or f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')} — {len(data.filas)} registro(s)"
    ths = "".join(f"<th>{escape(col.label)}</th>" for col in data.columnas)
    filas_html = []
    for fila in data.filas:
        tds = "".join(f"<td>{escape(_fmt(fila.get(col.key)))}</td>" for col in data.columnas)
        filas_html.append(f"<tr>{tds}</tr>")
    cuerpo = "".join(filas_html) or f'<tr><td colspan="{len(data.columnas)}" class="vacio">Sin registros</td></tr>'

    return f"""<!DOCTYPE html>
<html lang="es"><head><meta charset="utf-8"><style>
  @page {{ size: A4 landscape; margin: 1.2cm; }}
  * {{ font-family: 'Helvetica Neue', Arial, sans-serif; }}
  body {{ color: #1f2d3d; }}
  .titulo {{ font-size: 18px; font-weight: 700; color: #0e1c34; margin: 0 0 2px; }}
  .sub {{ font-size: 10px; color: #6a7e96; margin: 0 0 14px; }}
  table {{ width: 100%; border-collapse: collapse; }}
  thead th {{ background: #0e1c34; color: #fff; font-size: 9px; text-align: left;
              padding: 7px 8px; text-transform: uppercase; letter-spacing: .4px; }}
  tbody td {{ font-size: 9px; padding: 6px 8px; border-bottom: 1px solid #e3e9f1; }}
  tbody tr:nth-child(even) {{ background: #f5f8fc; }}
  .vacio {{ text-align: center; color: #98a8bc; padding: 18px; }}
</style></head><body>
  <div class="titulo">{escape(data.titulo)}</div>
  <div class="sub">{escape(sub)}</div>
  <table><thead><tr>{ths}</tr></thead><tbody>{cuerpo}</tbody></table>
</body></html>"""


@router.post("/pdf")
def export_pdf(data: ExportRequest, current_user = Depends(get_current_user), _=Depends(_limitar_body)):
    _validar_filas(data)
    from weasyprint import HTML
    pdf_bytes = HTML(string=_build_html(data)).write_pdf()
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{_nombre_archivo(data.titulo, "pdf")}"'},
    )
