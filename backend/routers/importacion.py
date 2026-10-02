"""
Importación masiva de ACTIVOS desde Excel (.xlsx).

Diseño centrado en SEGURIDAD:
  - Flujo en dos pasos: /validar (NO toca la BD) → preview → /ejecutar (re-valida).
  - Import parcial: cada fila se importa o se omite por completo (savepoint por fila),
    nunca se crea un activo a medias.
  - Nunca sobrescribe datos: placa/serial duplicados → error de fila.
  - Respeta la placa del Excel y luego ajusta el consecutivo para que las placas
    futuras generadas por la UI no colisionen.

Endpoints (prefix /api/importacion), todos gated por 'importacion.ejecutar':
  GET  /plantilla/activos          → descarga la plantilla .xlsx
  POST /activos/validar            → reporte de validación (sin escribir)
  POST /activos/ejecutar           → importa filas válidas/advertencia, reporte
  POST /activos/exportar-fallidas  → .xlsx con solo las filas con error + Motivo
"""
import io
import re
import unicodedata
from datetime import date, datetime

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, Body
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from database import get_db
from models.activo import Activo
from models.accesorio import Accesorio
from models.empresa import Empresa
from models.usuario import Usuario
from models.catalogo import Catalogo
from models.consecutivo import Consecutivo
from models.historial import HistorialMovimiento
from dependencies.rbac import require_permission
from services.rbac_service import get_user_empresa_ids, empresas_pueden_compartir
from services.inventario_service import calcular_fecha_obsolescencia, calcular_garantia_fin, resolver_ubicacion_catalogo

router = APIRouter(prefix="/api/importacion", tags=["Importación"])

# ── Tope de tamaño de subida ──────────────────────────────────────────────────
# 20 MB cubre ~50k filas holgadamente; los imports reales son de miles de filas
# (<5 MB), así que no rechaza archivos legítimos. Lee por bloques y rechaza ANTES
# de materializar el archivo completo en memoria (DoS por subida gigante).
MAX_UPLOAD_BYTES = 20 * 1024 * 1024  # 20 MB


async def _leer_upload_capado(file: UploadFile, max_bytes: int = MAX_UPLOAD_BYTES) -> bytes:
    chunks, total = [], 0
    while True:
        chunk = await file.read(1024 * 1024)  # 1 MB
        if not chunk:
            break
        total += len(chunk)
        if total > max_bytes:
            raise HTTPException(
                status_code=413,
                detail=f"El archivo supera el máximo de {max_bytes // (1024 * 1024)} MB",
            )
        chunks.append(chunk)
    return b"".join(chunks)

# ── Columnas de la plantilla (orden y nombres canónicos) ──────────────────────
COLUMNAS = [
    "Placa", "Empresa", "Tipo", "Estado", "Cédula", "Marca", "Modelo", "Serial",
    "Numero de parte", "Procesador", "Memoria RAM", "Disco 1", "Disco 2",
    "Ubicacion", "Fecha de compra", "Garantia (meses)", "Observacion",
    "Usuario (referencia)",
]
EJEMPLO = [
    "SC0001", "Socia BPO", "Portatil", "asignado", "1045892310", "HP", "ProBook 445 G10",
    "5CD2394KLM", "ABC-123", "Ryzen 7", "16GB", "512GB SSD", "", "Sede Principal",
    "2026-01-15", "12", "Equipo de ejemplo — borrar esta fila", "Juan Pérez",
]

ESTADOS_RECURSO = [
    "disponible", "asignado", "mantenimiento_preventivo", "mantenimiento_correctivo",
    "en_reparacion", "en_garantia", "retirado", "reservado",
]

# Mapa flexible de Estado (texto normalizado → canónico)
ESTADO_MAP = {
    "disponible": "disponible", "en bodega": "disponible", "bodega": "disponible",
    "asignado": "asignado", "en uso": "asignado", "entregado": "asignado",
    "mantenimiento": "mantenimiento_correctivo", "en mantenimiento": "mantenimiento_correctivo",
    "mantenimiento correctivo": "mantenimiento_correctivo",
    "mantenimiento preventivo": "mantenimiento_preventivo",
    "garantia": "en_garantia", "en garantia": "en_garantia",
    "reparacion": "en_reparacion", "en reparacion": "en_reparacion",
    "donde proveedor": "en_reparacion", "donde fabricante": "en_reparacion",
    "retirado": "retirado", "baja": "retirado", "dado de baja": "retirado",
    "reservado": "reservado",
}


# ── Helpers de normalización ──────────────────────────────────────────────────
def _norm(s) -> str:
    """minúsculas, sin acentos, espacios colapsados, recortado."""
    if s is None:
        return ""
    s = str(s).strip().lower()
    s = "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s)


def _cell_str(v) -> str:
    """Valor de celda → string limpio (para campos de texto y datos_raw)."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()


def _parse_fecha(v):
    """Devuelve (date|None, error_str|None). Vacío → (None, None) = sin fecha."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None, None
    if isinstance(v, datetime):
        return v.date(), None
    if isinstance(v, date):
        return v, None
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date(), None
        except ValueError:
            continue
    return None, f"Fecha de compra '{s}' no tiene un formato válido (use AAAA-MM-DD o DD/MM/AAAA)"


# ── Núcleo de validación (compartido por /validar y /ejecutar; NO escribe) ────
def _validar_archivo(db: Session, current_user, contenido: bytes) -> dict:
    """Parsea y valida cada fila. NO toca la BD para escribir (solo lecturas).
    Devuelve el reporte completo + datos normalizados por fila."""
    try:
        wb = load_workbook(filename=io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="No se pudo leer el archivo. ¿Es un .xlsx válido?")
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="El archivo está vacío")

    # Mapear encabezados (fila 1) → índice de columna, por nombre normalizado
    headers = rows[0]
    col_idx = {}
    for i, h in enumerate(headers):
        key = _norm(h)
        if key:
            col_idx[key] = i
    requeridas = ["placa", "empresa", "tipo"]
    faltan = [c for c in requeridas if c not in col_idx]
    if faltan:
        raise HTTPException(status_code=400,
                            detail=f"Faltan columnas obligatorias en la plantilla: {', '.join(faltan)}")

    def cell(row, header_norm):
        i = col_idx.get(header_norm)
        return row[i] if (i is not None and i < len(row)) else None

    # ── Catálogos / catálogos vivos (lecturas, en memoria) ──
    empresas = db.query(Empresa).all()
    emp_por_nombre = {_norm(e.nombre_empresa): e for e in empresas}
    empresa_ids_user = get_user_empresa_ids(db, current_user.id)  # None = super_admin (todas)

    tipos_cat = db.query(Catalogo).filter(
        Catalogo.categoria == "tipo_activo", Catalogo.activo == True).all()
    tipo_por_norm = {_norm(c.valor): c.valor for c in tipos_cat}

    # Placas y seriales existentes en BD (para detectar duplicados)
    placas_db = {(_cell_str(p[0]).upper()) for p in db.query(Activo.id_placa_activo).all()}
    seriales_db = {(_cell_str(s[0]).lower()) for s in
                   db.query(Activo.serial).filter(Activo.serial.isnot(None)).all() if _cell_str(s[0])}

    # Acumuladores intra-archivo (para duplicados dentro del mismo Excel)
    placas_file = {}    # placa_upper → fila_num
    seriales_file = {}  # serial_lower → fila_num

    filas = []
    resumen_errores = []
    total = validas = con_adv = con_err = 0

    for ridx, row in enumerate(rows[1:], start=2):
        # Saltar filas completamente vacías
        if row is None or all((c is None or str(c).strip() == "") for c in row):
            continue
        # Saltar la fila de ejemplo si quedó
        placa_raw = _cell_str(cell(row, "placa"))
        if "ejemplo" in _norm(placa_raw) or "borrar" in _norm(_cell_str(cell(row, "observacion"))):
            continue

        total += 1
        errores, advertencias = [], []
        datos = {}
        # Conservar valores crudos por columna (para el reporte de fallidas)
        datos_raw = {h: _cell_str(cell(row, _norm(h))) for h in COLUMNAS}

        # ── Placa (requerida, única) ──
        placa = placa_raw
        if not placa:
            errores.append("Placa es obligatoria")
        else:
            pu = placa.upper()
            if pu in placas_db:
                errores.append(f"La placa '{placa}' ya existe en el sistema")
            if pu in placas_file:
                errores.append(f"La placa '{placa}' está duplicada en el archivo (fila {placas_file[pu]})")
            placas_file.setdefault(pu, ridx)
        datos["id_placa_activo"] = placa

        # ── Empresa (requerida, match por nombre, con control de acceso) ──
        emp_raw = _cell_str(cell(row, "empresa"))
        empresa = emp_por_nombre.get(_norm(emp_raw)) if emp_raw else None
        if not emp_raw:
            errores.append("Empresa es obligatoria")
        elif not empresa:
            errores.append(f"Empresa '{emp_raw}' no encontrada")
        elif empresa_ids_user is not None and empresa.id not in empresa_ids_user:
            errores.append(f"No tienes acceso a la empresa '{empresa.nombre_empresa}'")
        if empresa:
            datos["empresa_id"] = empresa.id
            datos["_empresa_nombre"] = empresa.nombre_empresa
            datos["_empresa_prefijo"] = empresa.prefijo

        # ── Tipo (requerido, contra catálogo) ──
        tipo_raw = _cell_str(cell(row, "tipo"))
        tipo_canon = tipo_por_norm.get(_norm(tipo_raw)) if tipo_raw else None
        if not tipo_raw:
            errores.append("Tipo es obligatorio")
        elif not tipo_canon:
            errores.append(f"Tipo '{tipo_raw}' no existe en el catálogo")
        datos["tipo_activo"] = tipo_canon

        # ── Cédula (cruda; se valida junto con Estado) ──
        cedula = _cell_str(cell(row, "cedula"))

        # ── Estado (normalizar; inferir si vacío) ──
        estado_raw = _cell_str(cell(row, "estado"))
        if not estado_raw:
            estado = "asignado" if cedula else "disponible"
        else:
            estado = ESTADO_MAP.get(_norm(estado_raw))
            if not estado:
                errores.append(f"Estado '{estado_raw}' no es válido")
        datos["estado"] = estado

        # ── Cédula → asignación (si estado=asignado) ──
        datos["id_usuario"] = None
        datos["_cedula"] = cedula
        if estado == "asignado":
            if not cedula:
                errores.append("Cédula es obligatoria cuando el estado es 'asignado'")
            elif empresa:
                empleados = db.query(Usuario).filter(
                    Usuario.documento == cedula, Usuario.estado == "activo").all()
                elegible = next(
                    (u for u in empleados if empresas_pueden_compartir(db, empresa.id, u.empresa_id)), None)
                if not empleados:
                    errores.append(f"No existe un empleado activo con cédula '{cedula}'")
                elif not elegible:
                    errores.append(
                        f"El empleado con cédula '{cedula}' no pertenece a '{empresa.nombre_empresa}' ni a una empresa relacionada")
                else:
                    datos["id_usuario"] = elegible.id
                    datos["_empleado_nombre"] = elegible.nombre_completo

        # ── Serial (opcional, único si está presente) ──
        serial = _cell_str(cell(row, "serial"))
        if serial:
            sl = serial.lower()
            if sl in seriales_db:
                errores.append(f"El serial '{serial}' ya existe en el sistema")
            if sl in seriales_file:
                errores.append(f"El serial '{serial}' está duplicado en el archivo (fila {seriales_file[sl]})")
            seriales_file.setdefault(sl, ridx)
        datos["serial"] = serial or None

        # ── Texto passthrough ──
        for campo, header in [("marca", "marca"), ("modelo", "modelo"),
                              ("numero_parte", "numero de parte"), ("procesador", "procesador"),
                              ("memoria_ram", "memoria ram"), ("disco_1", "disco 1"),
                              ("disco_2", "disco 2"),
                              ("observaciones", "observacion")]:
            val = _cell_str(cell(row, header))
            datos[campo] = val or None

        # ── Ubicación: si se da, debe existir en el catálogo de la empresa (canónica) ──
        ubic_raw = _cell_str(cell(row, "ubicacion"))
        if ubic_raw and empresa:
            canon, ok = resolver_ubicacion_catalogo(db, empresa.id, ubic_raw)
            if not ok:
                errores.append(f"La ubicación '{ubic_raw}' no existe en el catálogo de la empresa '{empresa.nombre_empresa}'. Créala en Catálogos.")
            datos["ubicacion"] = canon if ok else None
        else:
            datos["ubicacion"] = ubic_raw or None

        # ── Fecha de compra (opcional; vacío=advertencia) ──
        fc, ferr = _parse_fecha(cell(row, "fecha de compra"))
        if ferr:
            errores.append(ferr)
        elif fc is None:
            advertencias.append("Sin fecha de compra: no se calculará obsolescencia ni garantía")
        datos["fecha_compra"] = fc.isoformat() if fc else None

        # ── Garantía (meses, opcional, entero ≥ 0) ──
        gm_raw = cell(row, "garantia (meses)")
        gm = None
        if gm_raw is not None and str(gm_raw).strip() != "":
            try:
                gm = int(float(str(gm_raw).strip()))
                if gm < 0:
                    errores.append("Garantía (meses) no puede ser negativa")
                    gm = None
            except (ValueError, TypeError):
                errores.append(f"Garantía (meses) '{gm_raw}' no es un número entero")
        datos["garantia_meses"] = gm

        # ── Resultado de la fila ──
        if errores:
            estado_res = "error"; con_err += 1
            for m in errores:
                resumen_errores.append(f"Fila {ridx}: {m}")
        elif advertencias:
            estado_res = "advertencia"; con_adv += 1
        else:
            estado_res = "valida"; validas += 1

        filas.append({
            "fila_num": ridx,
            "placa": placa,
            "estado_resultado": estado_res,
            "mensajes": errores + advertencias,
            "datos_normalizados": datos,
            "datos_raw": datos_raw,
        })

    wb.close()
    return {
        "total_filas": total,
        "validas": validas,
        "con_advertencia": con_adv,
        "con_error": con_err,
        "filas": filas,
        "resumen_errores": resumen_errores,
    }


# ── GET /plantilla/activos ────────────────────────────────────────────────────
@router.get("/plantilla/activos")
def descargar_plantilla(
    db: Session = Depends(get_db),
    current_user = require_permission("importacion.ejecutar"),
):
    wb = Workbook()
    ws = wb.active
    ws.title = "Activos"
    hdr_fill = PatternFill("solid", fgColor="1E3A5F")
    hdr_font = Font(color="FFFFFF", bold=True)
    for i, h in enumerate(COLUMNAS, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = hdr_fill; c.font = hdr_font
    # Fila de ejemplo (marcada para borrar)
    for i, v in enumerate(EJEMPLO, start=1):
        ws.cell(row=2, column=i, value=v).font = Font(italic=True, color="999999")
    for i in range(1, len(COLUMNAS) + 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 20

    # Hoja Instrucciones
    ins = wb.create_sheet("Instrucciones")
    lineas = [
        "IMPORTACIÓN DE ACTIVOS — INSTRUCCIONES",
        "",
        "1. Borra la fila de EJEMPLO antes de cargar.",
        "2. Columnas OBLIGATORIAS: Placa, Empresa, Tipo.",
        "   - Cédula es obligatoria SOLO si el Estado es 'asignado'.",
        "3. Placa: debe ser ÚNICA (no puede repetirse con activos existentes ni dentro del archivo).",
        "4. Empresa: escribe el nombre tal como está registrado en el sistema.",
        "5. Tipo: debe coincidir con un tipo del catálogo (ver hoja 'Valores válidos').",
        "6. Estado: usa uno de los valores válidos (ver hoja 'Valores válidos').",
        "   - Si lo dejas vacío: será 'asignado' si pones Cédula, o 'disponible' si no.",
        "7. Cédula: el empleado debe existir, estar activo y pertenecer a la empresa o a una empresa relacionada.",
        "8. Serial: opcional; si lo pones, no puede repetirse con uno existente ni dentro del archivo.",
        "9. Fecha de compra: formato AAAA-MM-DD o DD/MM/AAAA. Si la dejas vacía, no se calculará obsolescencia/garantía (advertencia).",
        "10. Garantía (meses): número entero. Calcula automáticamente el fin de garantía desde la fecha de compra.",
        "11. La obsolescencia se calcula automáticamente desde la fecha de compra y el tipo.",
        "12. Ubicacion: opcional. Si la indicas, debe existir en el catálogo de Ubicaciones de la empresa de la fila",
        "    (Administración → Catálogos → Ubicación). Si no existe, la fila se rechaza.",
        "13. 'Usuario (referencia)' es solo informativo; la asignación se hace por la Cédula.",
        "",
        "El sistema primero VALIDA (vista previa) y NO importa nada hasta que confirmes.",
        "Se importan solo las filas válidas; las filas con error se omiten y puedes descargarlas para corregir.",
        "IMPORTANTE: haz un respaldo de la base de datos antes de importar.",
    ]
    for r, t in enumerate(lineas, start=1):
        cell = ins.cell(row=r, column=1, value=t)
        if r == 1:
            cell.font = Font(bold=True, size=14)
    ins.column_dimensions["A"].width = 110

    # Hoja Valores válidos (tipos del catálogo + estados)
    val = wb.create_sheet("Valores válidos")
    val.cell(row=1, column=1, value="Tipos de activo válidos (catálogo)").font = Font(bold=True)
    tipos = db.query(Catalogo.valor).filter(
        Catalogo.categoria == "tipo_activo", Catalogo.activo == True).order_by(Catalogo.valor).all()
    r = 2
    for (t,) in tipos:
        val.cell(row=r, column=1, value=t); r += 1
    val.cell(row=1, column=3, value="Estados válidos (texto aceptado)").font = Font(bold=True)
    estados_help = [
        "disponible (o: en bodega, bodega)",
        "asignado (o: en uso, entregado)",
        "mantenimiento (→ mantenimiento_correctivo)",
        "garantia / en garantia (→ en_garantia)",
        "reparacion / donde proveedor / donde fabricante (→ en_reparacion)",
        "retirado / baja / dado de baja",
        "reservado",
    ]
    for i, e in enumerate(estados_help, start=2):
        val.cell(row=i, column=3, value=e)
    val.column_dimensions["A"].width = 30
    val.column_dimensions["C"].width = 55

    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="plantilla_activos.xlsx"'},
    )


# ── POST /activos/validar (NO escribe) ────────────────────────────────────────
@router.post("/activos/validar")
async def validar_activos(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user = require_permission("importacion.ejecutar"),
):
    contenido = await _leer_upload_capado(file)
    reporte = _validar_archivo(db, current_user, contenido)
    # No exponer estructuras internas innecesarias al cliente, pero sí lo útil
    return reporte


# ── POST /activos/ejecutar (re-valida + importa parcial) ──────────────────────
def _suffix_de_placa(placa: str, prefijo: str):
    """Devuelve el entero del sufijo si la placa = prefijo + dígitos; si no, None."""
    m = re.fullmatch(re.escape(prefijo.upper()) + r"(\d+)", (placa or "").upper())
    return int(m.group(1)) if m else None


@router.post("/activos/ejecutar")
async def ejecutar_activos(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user = require_permission("importacion.ejecutar"),
):
    contenido = await _leer_upload_capado(file)
    # 1. Re-validar SIEMPRE en el servidor (nunca confiar en el cliente)
    reporte = _validar_archivo(db, current_user, contenido)

    importables = [f for f in reporte["filas"] if f["estado_resultado"] in ("valida", "advertencia")]
    fallidas = [f for f in reporte["filas"] if f["estado_resultado"] == "error"]

    importados = 0
    detalle = []
    # max sufijo de placa importada por empresa (para ajustar el consecutivo)
    max_suffix_por_empresa = {}   # empresa_id → (prefijo, max_int)

    for f in importables:
        d = f["datos_normalizados"]
        try:
            with db.begin_nested():   # savepoint por fila: o entra completa o no entra
                asignado = d["estado"] == "asignado" and d.get("id_usuario")
                fc = date.fromisoformat(d["fecha_compra"]) if d.get("fecha_compra") else None
                activo = Activo(
                    id_placa_activo=d["id_placa_activo"],
                    empresa_id=d["empresa_id"],
                    tipo_activo=d["tipo_activo"],
                    estado=d["estado"],
                    id_usuario=d.get("id_usuario"),
                    marca=d.get("marca"), modelo=d.get("modelo"), serial=d.get("serial"),
                    numero_parte=d.get("numero_parte"), procesador=d.get("procesador"),
                    memoria_ram=d.get("memoria_ram"), disco_1=d.get("disco_1"),
                    disco_2=d.get("disco_2"), ubicacion=d.get("ubicacion"),
                    observaciones=d.get("observaciones"),
                    fecha_compra=fc, garantia_meses=d.get("garantia_meses"),
                )
                # Auto-cálculos (la plantilla no trae valores manuales → siempre auto)
                _, fobs = calcular_fecha_obsolescencia(db, d["tipo_activo"], fc)
                if fobs:
                    activo.fecha_obsolescencia = fobs
                if d.get("garantia_meses") is not None:
                    activo.garantia_fin = calcular_garantia_fin(d["garantia_meses"], fc)
                db.add(activo)
                db.flush()
                # Historial (migración: SIN acta)
                if asignado:
                    obs = f"Importación masiva — asignado a {d.get('_cedula')} (sin acta)"
                    tipo_mov = "asignacion"
                else:
                    obs = "Importación masiva"
                    tipo_mov = "creacion"
                db.add(HistorialMovimiento(
                    id_activo=activo.id, tipo_movimiento=tipo_mov,
                    responsable=current_user.email, observaciones=obs))
            importados += 1
            detalle.append({"fila_num": f["fila_num"], "placa": d["id_placa_activo"], "resultado": "importado"})
            # Acumular sufijo para el bump del consecutivo
            pref = d.get("_empresa_prefijo")
            if pref:
                suf = _suffix_de_placa(d["id_placa_activo"], pref)
                if suf is not None:
                    prev = max_suffix_por_empresa.get(d["empresa_id"])
                    if prev is None or suf > prev[1]:
                        max_suffix_por_empresa[d["empresa_id"]] = (pref, suf)
        except Exception as e:
            # La fila falló al insertar (p. ej. choque de unicidad concurrente) → se omite, no a medias
            detalle.append({"fila_num": f["fila_num"], "placa": d.get("id_placa_activo"),
                            "resultado": "omitido", "motivo": str(e)[:200]})

    # 4. Ajustar consecutivos: ultimo_numero >= max sufijo importado (por empresa)
    for empresa_id, (pref, max_suf) in max_suffix_por_empresa.items():
        cons = db.query(Consecutivo).filter(
            Consecutivo.empresa_id == empresa_id, Consecutivo.tipo == "ACTIVO"
        ).with_for_update().first()
        if not cons:
            cons = Consecutivo(empresa_id=empresa_id, tipo="ACTIVO", ultimo_numero=0)
            db.add(cons)
        if max_suf > (cons.ultimo_numero or 0):
            cons.ultimo_numero = max_suf

    db.commit()

    return {
        "importados": importados,
        "omitidos": len(fallidas),
        "advertencias": reporte["con_advertencia"],
        "total_filas": reporte["total_filas"],
        "detalle": detalle,
        # filas con error (datos crudos + motivo) para descargar y corregir
        "filas_fallidas": [
            {"datos_raw": f["datos_raw"], "motivo": "; ".join(f["mensajes"])}
            for f in fallidas
        ],
    }


# ── POST /activos/exportar-fallidas → .xlsx solo de filas con error ──────────
@router.post("/activos/exportar-fallidas")
def exportar_fallidas(
    payload: dict = Body(...),
    current_user = require_permission("importacion.ejecutar"),
):
    filas = payload.get("filas_fallidas", [])
    wb = Workbook(); ws = wb.active; ws.title = "Filas con error"
    cols = COLUMNAS + ["Motivo"]
    hdr_fill = PatternFill("solid", fgColor="8B1A1A"); hdr_font = Font(color="FFFFFF", bold=True)
    for i, h in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=h); c.fill = hdr_fill; c.font = hdr_font
    for r, f in enumerate(filas, start=2):
        raw = f.get("datos_raw", {})
        for i, h in enumerate(COLUMNAS, start=1):
            ws.cell(row=r, column=i, value=raw.get(h, ""))
        ws.cell(row=r, column=len(cols), value=f.get("motivo", ""))
    for i in range(1, len(cols) + 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 22
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="filas_fallidas.xlsx"'},
    )


# ══════════════════════════════════════════════════════════════════════════════
#  ACCESORIOS — empresa elegida UNA vez en la UI (no es columna), más simple
# ══════════════════════════════════════════════════════════════════════════════
COLUMNAS_ACC = [
    "Placa", "Tipo", "Estado", "Cédula", "Marca", "Modelo", "No.Parte",
    "Serial", "Ubicacion", "Usuario (referencia)",
]
EJEMPLO_ACC = [
    "SCA0001", "Mouse", "asignado", "1045892310", "Logitech", "M170",
    "ABC-123", "SN-MOUSE-001", "Sede Principal", "Juan Pérez — borrar esta fila",
]


def _suffix_accesorio(placa: str, prefijo: str):
    """Sufijo entero si la placa = prefijo + 'A' + dígitos (formato SCA0001); si no, None."""
    m = re.fullmatch(re.escape(prefijo.upper()) + r"A(\d+)", (placa or "").upper())
    return int(m.group(1)) if m else None


def _validar_archivo_accesorios(db: Session, current_user, contenido: bytes, empresa_id: str) -> dict:
    """Valida un .xlsx de accesorios contra UNA empresa fija (no es columna).
    Solo lecturas a la BD. Mismo formato de reporte que activos."""
    empresa = db.query(Empresa).filter(Empresa.id == empresa_id).first()
    if not empresa:
        raise HTTPException(status_code=400, detail="Empresa no encontrada")
    empresa_ids_user = get_user_empresa_ids(db, current_user.id)
    if empresa_ids_user is not None and empresa_id not in empresa_ids_user:
        raise HTTPException(status_code=403, detail="No tienes acceso a esa empresa")

    try:
        wb = load_workbook(filename=io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="No se pudo leer el archivo. ¿Es un .xlsx válido?")
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="El archivo está vacío")

    col_idx = {}
    for i, h in enumerate(rows[0]):
        key = _norm(h)
        if key:
            col_idx[key] = i
    faltan = [c for c in ["placa", "tipo"] if c not in col_idx]
    if faltan:
        raise HTTPException(status_code=400,
                            detail=f"Faltan columnas obligatorias en la plantilla: {', '.join(faltan)}")

    def cell(row, header_norm):
        i = col_idx.get(header_norm)
        return row[i] if (i is not None and i < len(row)) else None

    tipos_cat = db.query(Catalogo).filter(
        Catalogo.categoria == "tipo_accesorio", Catalogo.activo == True).all()
    tipo_por_norm = {_norm(c.valor): c.valor for c in tipos_cat}

    placas_db = {(_cell_str(p[0]).upper()) for p in db.query(Accesorio.id_placa_accesorio).all()}
    seriales_db = {(_cell_str(s[0]).lower()) for s in
                   db.query(Accesorio.serial).filter(Accesorio.serial.isnot(None)).all() if _cell_str(s[0])}

    placas_file, seriales_file = {}, {}
    filas, resumen_errores = [], []
    total = validas = con_adv = con_err = 0

    for ridx, row in enumerate(rows[1:], start=2):
        if row is None or all((c is None or str(c).strip() == "") for c in row):
            continue
        placa_raw = _cell_str(cell(row, "placa"))
        usuario_ref = _cell_str(cell(row, "usuario (referencia)"))
        if "ejemplo" in _norm(placa_raw) or "borrar" in _norm(usuario_ref):
            continue

        total += 1
        errores, advertencias = [], []
        datos = {}
        datos_raw = {h: _cell_str(cell(row, _norm(h))) for h in COLUMNAS_ACC}

        # ── Placa (requerida, única) ──
        placa = placa_raw
        if not placa:
            errores.append("Placa es obligatoria")
        else:
            pu = placa.upper()
            if pu in placas_db:
                errores.append(f"La placa '{placa}' ya existe en el sistema")
            if pu in placas_file:
                errores.append(f"La placa '{placa}' está duplicada en el archivo (fila {placas_file[pu]})")
            placas_file.setdefault(pu, ridx)
        datos["id_placa_accesorio"] = placa

        # ── Tipo (requerido, contra catálogo tipo_accesorio) ──
        tipo_raw = _cell_str(cell(row, "tipo"))
        tipo_canon = tipo_por_norm.get(_norm(tipo_raw)) if tipo_raw else None
        if not tipo_raw:
            errores.append("Tipo es obligatorio")
        elif not tipo_canon:
            errores.append(f"Tipo '{tipo_raw}' no existe en el catálogo de accesorios")
        datos["tipo_accesorio"] = tipo_canon

        # ── Cédula (cruda) + Estado (normalizar/inferir) ──
        cedula = _cell_str(cell(row, "cedula"))
        estado_raw = _cell_str(cell(row, "estado"))
        if not estado_raw:
            estado = "asignado" if cedula else "disponible"
        else:
            estado = ESTADO_MAP.get(_norm(estado_raw))
            if not estado:
                errores.append(f"Estado '{estado_raw}' no es válido")
        datos["estado"] = estado

        # ── Cédula → asignación (empresa fija o hermana) ──
        datos["id_usuario"] = None
        datos["_cedula"] = cedula
        if estado == "asignado":
            if not cedula:
                errores.append("Cédula es obligatoria cuando el estado es 'asignado'")
            else:
                empleados = db.query(Usuario).filter(
                    Usuario.documento == cedula, Usuario.estado == "activo").all()
                elegible = next(
                    (u for u in empleados if empresas_pueden_compartir(db, empresa_id, u.empresa_id)), None)
                if not empleados:
                    errores.append(f"No existe un empleado activo con cédula '{cedula}'")
                elif not elegible:
                    errores.append(
                        f"El empleado con cédula '{cedula}' no pertenece a '{empresa.nombre_empresa}' ni a una empresa relacionada")
                else:
                    datos["id_usuario"] = elegible.id

        # ── Serial (opcional, único) ──
        serial = _cell_str(cell(row, "serial"))
        if serial:
            sl = serial.lower()
            if sl in seriales_db:
                errores.append(f"El serial '{serial}' ya existe en el sistema")
            if sl in seriales_file:
                errores.append(f"El serial '{serial}' está duplicado en el archivo (fila {seriales_file[sl]})")
            seriales_file.setdefault(sl, ridx)
        datos["serial"] = serial or None

        # ── Texto: marca, modelo ──
        datos["marca"] = _cell_str(cell(row, "marca")) or None
        datos["modelo"] = _cell_str(cell(row, "modelo")) or None

        # ── Ubicacion: si se da, debe existir en el catálogo de la empresa (canónica) ──
        ubic_raw = _cell_str(cell(row, "ubicacion"))
        if ubic_raw:
            canon, ok = resolver_ubicacion_catalogo(db, empresa_id, ubic_raw)
            if not ok:
                errores.append(f"La ubicación '{ubic_raw}' no existe en el catálogo de la empresa '{empresa.nombre_empresa}'. Créala en Catálogos.")
            datos["ubicacion"] = canon if ok else None
        else:
            datos["ubicacion"] = None

        # ── No.Parte: accesorios no tienen campo → se conserva en observaciones ──
        no_parte = _cell_str(cell(row, "no.parte"))
        datos["observaciones"] = (f"No. Parte: {no_parte}" if no_parte else None)

        if errores:
            estado_res = "error"; con_err += 1
            for m in errores:
                resumen_errores.append(f"Fila {ridx}: {m}")
        elif advertencias:
            estado_res = "advertencia"; con_adv += 1
        else:
            estado_res = "valida"; validas += 1

        filas.append({
            "fila_num": ridx, "placa": placa, "estado_resultado": estado_res,
            "mensajes": errores + advertencias,
            "datos_normalizados": datos, "datos_raw": datos_raw,
        })

    wb.close()
    return {
        "total_filas": total, "validas": validas, "con_advertencia": con_adv,
        "con_error": con_err, "filas": filas, "resumen_errores": resumen_errores,
        "_empresa_prefijo": empresa.prefijo, "_empresa_nombre": empresa.nombre_empresa,
    }


# ── GET /plantilla/accesorios ─────────────────────────────────────────────────
@router.get("/plantilla/accesorios")
def descargar_plantilla_accesorios(
    db: Session = Depends(get_db),
    current_user = require_permission("importacion.ejecutar"),
):
    wb = Workbook(); ws = wb.active; ws.title = "Accesorios"
    hdr_fill = PatternFill("solid", fgColor="1E3A5F"); hdr_font = Font(color="FFFFFF", bold=True)
    for i, h in enumerate(COLUMNAS_ACC, start=1):
        c = ws.cell(row=1, column=i, value=h); c.fill = hdr_fill; c.font = hdr_font
    for i, v in enumerate(EJEMPLO_ACC, start=1):
        ws.cell(row=2, column=i, value=v).font = Font(italic=True, color="999999")
    for i in range(1, len(COLUMNAS_ACC) + 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 20

    ins = wb.create_sheet("Instrucciones")
    lineas = [
        "IMPORTACIÓN DE ACCESORIOS — INSTRUCCIONES",
        "",
        "1. La EMPRESA se elige en la aplicación (NO es una columna): todas las filas",
        "   del archivo se importan a la empresa seleccionada.",
        "2. Borra la fila de EJEMPLO antes de cargar.",
        "3. Columnas OBLIGATORIAS: Placa, Tipo. (Cédula es obligatoria solo si Estado='asignado').",
        "4. Placa: debe ser ÚNICA (no puede repetirse con accesorios existentes ni dentro del archivo).",
        "5. Tipo: debe coincidir con un tipo del catálogo de accesorios (ver 'Valores válidos').",
        "6. Estado: usa un valor válido. 'BAJA' o 'dado de baja' se importan directamente como 'retirado'",
        "   (baja histórica, sin flujo de aprobación). Si lo dejas vacío: 'asignado' si pones Cédula, si no 'disponible'.",
        "7. Cédula: el empleado debe existir, estar activo y pertenecer a la empresa seleccionada o a una relacionada.",
        "8. Serial: opcional; si lo pones, no puede repetirse.",
        "9. Ubicacion: opcional. Si la indicas, debe existir en el catálogo de Ubicaciones de la empresa seleccionada",
        "   (Administración → Catálogos → Ubicación). Las ubicaciones válidas dependen de la empresa elegida. Si no existe, la fila se rechaza.",
        "10. No.Parte: se guarda dentro de Observaciones (los accesorios no tienen un campo de número de parte).",
        "11. 'Usuario (referencia)' es solo informativo; la asignación se hace por la Cédula.",
        "",
        "El sistema primero VALIDA (vista previa) y NO importa nada hasta que confirmes.",
        "IMPORTANTE: haz un respaldo de la base de datos antes de importar.",
    ]
    for r, t in enumerate(lineas, start=1):
        c = ins.cell(row=r, column=1, value=t)
        if r == 1:
            c.font = Font(bold=True, size=14)
    ins.column_dimensions["A"].width = 110

    val = wb.create_sheet("Valores válidos")
    val.cell(row=1, column=1, value="Tipos de accesorio válidos (catálogo)").font = Font(bold=True)
    tipos = db.query(Catalogo.valor).filter(
        Catalogo.categoria == "tipo_accesorio", Catalogo.activo == True).order_by(Catalogo.valor).all()
    r = 2
    for (t,) in tipos:
        val.cell(row=r, column=1, value=t); r += 1
    val.cell(row=1, column=3, value="Estados válidos (texto aceptado)").font = Font(bold=True)
    estados_help = [
        "disponible (o: en bodega, bodega)",
        "asignado (o: en uso, entregado)",
        "mantenimiento (→ mantenimiento_correctivo)",
        "garantia / en garantia (→ en_garantia)",
        "reparacion / donde proveedor / donde fabricante (→ en_reparacion)",
        "retirado / baja / dado de baja  (→ retirado, baja histórica)",
        "reservado",
    ]
    for i, e in enumerate(estados_help, start=2):
        val.cell(row=i, column=3, value=e)
    val.column_dimensions["A"].width = 30
    val.column_dimensions["C"].width = 60

    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="plantilla_accesorios.xlsx"'},
    )


# ── POST /accesorios/validar (NO escribe) ─────────────────────────────────────
@router.post("/accesorios/validar")
async def validar_accesorios(
    file: UploadFile = File(...),
    empresa_id: str = Form(...),
    db: Session = Depends(get_db),
    current_user = require_permission("importacion.ejecutar"),
):
    contenido = await _leer_upload_capado(file)
    return _validar_archivo_accesorios(db, current_user, contenido, empresa_id)


# ── POST /accesorios/ejecutar (re-valida + importa parcial) ───────────────────
@router.post("/accesorios/ejecutar")
async def ejecutar_accesorios(
    file: UploadFile = File(...),
    empresa_id: str = Form(...),
    db: Session = Depends(get_db),
    current_user = require_permission("importacion.ejecutar"),
):
    contenido = await _leer_upload_capado(file)
    reporte = _validar_archivo_accesorios(db, current_user, contenido, empresa_id)
    prefijo = reporte["_empresa_prefijo"]

    importables = [f for f in reporte["filas"] if f["estado_resultado"] in ("valida", "advertencia")]
    fallidas = [f for f in reporte["filas"] if f["estado_resultado"] == "error"]

    importados = 0
    detalle = []
    max_suffix = None

    for f in importables:
        d = f["datos_normalizados"]
        try:
            with db.begin_nested():   # savepoint por fila
                asignado = d["estado"] == "asignado" and d.get("id_usuario")
                acc = Accesorio(
                    id_placa_accesorio=d["id_placa_accesorio"],
                    empresa_id=empresa_id,
                    tipo_accesorio=d["tipo_accesorio"],
                    estado=d["estado"],
                    id_usuario=d.get("id_usuario"),
                    marca=d.get("marca"), modelo=d.get("modelo"), serial=d.get("serial"),
                    ubicacion=d.get("ubicacion"), observaciones=d.get("observaciones"),
                )
                db.add(acc); db.flush()
                if asignado:
                    obs = f"Importación masiva — asignado a {d.get('_cedula')} (sin acta)"
                    tipo_mov = "asignacion"
                elif d["estado"] == "retirado":
                    obs = "Importación masiva — baja histórica"
                    tipo_mov = "baja"
                else:
                    obs = "Importación masiva"
                    tipo_mov = "creacion"
                db.add(HistorialMovimiento(
                    id_accesorio=acc.id, tipo_movimiento=tipo_mov,
                    responsable=current_user.email, observaciones=obs))
            importados += 1
            detalle.append({"fila_num": f["fila_num"], "placa": d["id_placa_accesorio"], "resultado": "importado"})
            suf = _suffix_accesorio(d["id_placa_accesorio"], prefijo)
            if suf is not None and (max_suffix is None or suf > max_suffix):
                max_suffix = suf
        except Exception as e:
            detalle.append({"fila_num": f["fila_num"], "placa": d.get("id_placa_accesorio"),
                            "resultado": "omitido", "motivo": str(e)[:200]})

    # Bump del consecutivo tipo ACCESORIO de la empresa
    if max_suffix is not None:
        cons = db.query(Consecutivo).filter(
            Consecutivo.empresa_id == empresa_id, Consecutivo.tipo == "ACCESORIO"
        ).with_for_update().first()
        if not cons:
            cons = Consecutivo(empresa_id=empresa_id, tipo="ACCESORIO", ultimo_numero=0)
            db.add(cons)
        if max_suffix > (cons.ultimo_numero or 0):
            cons.ultimo_numero = max_suffix

    db.commit()

    return {
        "importados": importados,
        "omitidos": len(fallidas),
        "advertencias": reporte["con_advertencia"],
        "total_filas": reporte["total_filas"],
        "detalle": detalle,
        "filas_fallidas": [
            {"datos_raw": f["datos_raw"], "motivo": "; ".join(f["mensajes"])}
            for f in fallidas
        ],
    }


# ── POST /accesorios/exportar-fallidas → .xlsx solo de filas con error ────────
@router.post("/accesorios/exportar-fallidas")
def exportar_fallidas_accesorios(
    payload: dict = Body(...),
    current_user = require_permission("importacion.ejecutar"),
):
    filas = payload.get("filas_fallidas", [])
    wb = Workbook(); ws = wb.active; ws.title = "Filas con error"
    cols = COLUMNAS_ACC + ["Motivo"]
    hdr_fill = PatternFill("solid", fgColor="8B1A1A"); hdr_font = Font(color="FFFFFF", bold=True)
    for i, h in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=h); c.fill = hdr_fill; c.font = hdr_font
    for r, f in enumerate(filas, start=2):
        raw = f.get("datos_raw", {})
        for i, h in enumerate(COLUMNAS_ACC, start=1):
            ws.cell(row=r, column=i, value=raw.get(h, ""))
        ws.cell(row=r, column=len(cols), value=f.get("motivo", ""))
    for i in range(1, len(cols) + 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 22
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="filas_fallidas_accesorios.xlsx"'},
    )
